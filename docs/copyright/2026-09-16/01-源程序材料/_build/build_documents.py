from __future__ import annotations

import ast
import csv
import hashlib
import io
import json
import math
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
DATA = json.loads((HERE / 'source-data.json').read_text())
ROOT = Path(DATA['root'])
TITLE = '心语大学生心理健康平台'
VERSION = 'V1.0'
CODE_NAME = f'{TITLE}_{VERSION}_源程序_前30页后30页.docx'


def code_flags(record):
    lines = record['lines']
    flags = [bool(line.strip()) for line in lines]
    if record['path'].endswith('.py'):
        tree = ast.parse('\n'.join(lines))
        docstrings = set()
        for node in ast.walk(tree):
            body = getattr(node, 'body', None)
            if isinstance(body, list) and body and isinstance(body[0], ast.Expr):
                value = body[0].value
                if isinstance(value, ast.Constant) and isinstance(value.value, str):
                    docstrings.update(range(body[0].lineno, body[0].end_lineno + 1))
        for i, line in enumerate(lines):
            if line.lstrip().startswith('#') or i + 1 in docstrings:
                flags[i] = False
    else:
        in_comment = False
        for i, line in enumerate(lines):
            s = line.strip()
            if in_comment:
                flags[i] = False
                if '*/' in s or '-->' in s:
                    in_comment = False
                continue
            if s.startswith('//') or s.startswith('#'):
                flags[i] = False
            elif s.startswith('/*') or s.startswith('<!--'):
                flags[i] = False
                in_comment = not ('*/' in s or '-->' in s)
    return flags


stream = []
for record in DATA['records']:
    flags = code_flags(record)
    record['counted_code_lines'] = sum(flags)
    for n, (line, flag) in enumerate(zip(record['lines'], flags), 1):
        stream.append({'file': record['path'], 'line': n, 'text': line, 'code': flag})

total_code = sum(r['code'] for r in stream)
# Distribute the remainder so the final page also contains at least 50 real code lines.
page_count = total_code // 50
extra = total_code % 50
targets = [50 + (i < extra) for i in range(page_count)]
all_pages, cursor = [], 0
for i, target in enumerate(targets):
    start, count = cursor, 0
    while cursor < len(stream) and count < target:
        count += stream[cursor]['code']
        cursor += 1
    if i == page_count - 1:
        cursor = len(stream)
    all_pages.append({'full_page': i + 1, 'code_lines': count, 'rows': stream[start:cursor]})
assert cursor == len(stream)
assert all(p['code_lines'] >= 50 for p in all_pages)
assert len(all_pages) > 60
selected = all_pages[:30] + all_pages[-30:]
assert selected[-1]['rows'][-1] == stream[-1]


def font(run, name='Arial', size=10, east='Arial Unicode MS', color=None, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.set(qn('w:eastAsia'), east)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def no_grid(paragraph):
    e = OxmlElement('w:snapToGrid')
    e.set(qn('w:val'), '0')
    paragraph._p.get_or_add_pPr().append(e)


def field(paragraph, name, cached):
    run = paragraph.add_run()
    font(run, size=8)
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    inst = OxmlElement('w:instrText')
    inst.text = name
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    value = OxmlElement('w:t')
    value.text = str(cached)
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for e in [begin, inst, separate, value, end]:
        run._r.append(e)


def setup(doc, code=False):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(1.7), Cm(1.5)
    s.left_margin = s.right_margin = Cm(1.25 if code else 2)
    s.header_distance, s.footer_distance = Cm(.65), Cm(.65)
    st = doc.styles['Normal']
    st.font.name, st.font.size = 'Arial', Pt(10)
    st.paragraph_format.space_after = Pt(0 if code else 7)
    header = s.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font(header.add_run(f'{TITLE}  {VERSION}'), size=10, bold=True)
    border = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for k, v in {'val':'single','sz':'4','color':'999999','space':'5'}.items():
        bottom.set(qn('w:' + k), v)
    border.append(bottom)
    header._p.get_or_add_pPr().append(border)
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('源程序鉴别材料  ·  第 ' if code else '源程序材料核对说明  ·  第 '), size=8)
    field(footer, 'PAGE', 1)
    font(footer.add_run(' 页'), size=8)
    if code:
        font(footer.add_run(' / 共 60 页'), size=8)
    doc.core_properties.title = f'{TITLE} {VERSION} ' + ('源程序鉴别材料' if code else '代码材料核对说明与模块索引')
    doc.core_properties.author = ''
    doc.core_properties.subject = '源程序前、后各连续30页' if code else '源代码统计、追溯和功能模块对应'
    doc.core_properties.comments = ''


doc = Document()
setup(doc, code=True)
manifest_pages = []
for submit_no, page in enumerate(selected, 1):
    # Small source labels and blank lines are not counted toward the 50 code lines.
    h = doc.add_paragraph()
    h.paragraph_format.page_break_before = submit_no > 1
    h.paragraph_format.space_after = Pt(5)
    h.paragraph_format.line_spacing = Pt(11)
    h.paragraph_format.keep_with_next = True
    side = '前部' if submit_no <= 30 else '后部'
    font(h.add_run(f'{side}连续30页  |  全文第 {page["full_page"]} 页  |  本页 {page["code_lines"]} 行代码'), size=8, color='555555')
    current_file = None
    ranges = []
    nonblank_rows = 0
    # Stable code typography. Word may wrap long string literals without changing the stored source.
    for row in page['rows']:
        if row['file'] != current_file:
            current_file = row['file']
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = Pt(10)
            p.paragraph_format.keep_with_next = True
            font(p.add_run(current_file), name='Arial', size=8, bold=True)
            ranges.append({'file': current_file, 'start_line': row['line'], 'end_line': row['line']})
        else:
            ranges[-1]['end_line'] = row['line']
        p = doc.add_paragraph()
        f = p.paragraph_format
        f.space_before = f.space_after = Pt(0)
        f.widow_control = False
        f.keep_together = True
        no_grid(p)
        if not row['text'].strip():
            f.line_spacing = Pt(3)
            font(p.add_run(''), size=2)
            continue
        nonblank_rows += 1
        f.left_indent = Pt(29)
        f.first_line_indent = Pt(-29)
        f.tab_stops.add_tab_stop(Pt(29))
        f.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        f.line_spacing = Pt(11)
        font(p.add_run(f'{row["line"]:04d}\t'), name='Courier New', size=8, color='777777')
        font(p.add_run(row['text'].expandtabs(4)), name='Courier New', size=8.5)
    manifest_pages.append({'submission_page':submit_no, 'part':side, 'full_page':page['full_page'], 'code_lines':page['code_lines'], 'nonblank_source_rows':nonblank_rows, 'ranges':ranges})
doc.save(OUT / CODE_NAME)

# Exact source reconstruction check: row texts in the Word file must equal the frozen formatted source.
reopened = Document(OUT / CODE_NAME)
actual = [p.text.split('\t', 1)[1] for p in reopened.paragraphs if re.match(r'^\d{4,5}\t', p.text)]
expected = [r['text'].expandtabs(4) for p in selected for r in p['rows'] if r['text'].strip()]
assert actual == expected, 'Word source text does not match the frozen source'

summary = {
    'software_name': TITLE, 'registration_version': VERSION,
    'name_and_version_status': 'provisional pending applicant confirmation; distinct from repository and package versions',
    'source_root': str(ROOT), 'snapshot_time': DATA['created_at'], 'git_commit': DATA['commit'],
    'source_state': DATA['source'],
    'source_files': len(DATA['records']),
    'original_physical_lines': sum(r['original_lines'] for r in DATA['records']),
    'original_nonblank_lines': sum(r['original_nonblank'] for r in DATA['records']),
    'formatted_physical_lines': sum(r['formatted_lines'] for r in DATA['records']),
    'formatted_nonblank_lines': sum(r['formatted_nonblank'] for r in DATA['records']),
    'formatted_nonblank_noncomment_lines': total_code,
    'typescript_ast_equivalence_files': DATA['transpile_checks'],
    'full_listing_pages': page_count, 'submitted_pages': 60,
    'submitted_counted_code_lines': sum(p['code_lines'] for p in selected),
    'final_module': selected[-1]['rows'][-1]['file'],
    'final_module_last_line': selected[-1]['rows'][-1]['line'],
    'final_module_last_text': selected[-1]['rows'][-1]['text'],
    'word_text_roundtrip_verified': True,
    'pages': manifest_pages,
    'files': [{k:v for k,v in r.items() if k != 'lines'} for r in DATA['records']],
}
(HERE / 'manifest.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2))
with (OUT / '代码文件与行数清单.csv').open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(['编排顺序','文件路径','原始物理行数','原始非空行数','排版后物理行数','排版后非空非纯注释行数','原文件SHA256'])
    for i, r in enumerate(DATA['records'], 1):
        writer.writerow([i,r['path'],r['original_lines'],r['original_nonblank'],r['formatted_lines'],r['counted_code_lines'],r['original_sha256']])
with (HERE / '完整源程序编排.txt').open('w') as f:
    for page in all_pages:
        f.write(f'\n=== 全文第 {page["full_page"]} 页 / {page_count} 页 ===\n')
        previous = None
        for row in page['rows']:
            if row['file'] != previous:
                f.write(f'\n=== 文件：{row["file"]} ===\n')
                previous = row['file']
            f.write(f'{row["line"]:04d}\t{row["text"]}\n')


def para(d, text, bold=False, size=10):
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    font(p.add_run(text), size=size, bold=bold)
    return p


notes = Document()
setup(notes)
para(notes, '代码材料核对说明与模块索引', bold=True, size=18)
para(notes, '本文件用于核对和后续编写说明书，不计入60页源程序鉴别材料。', size=10)
para(notes, '1. 本次材料范围', bold=True, size=12)
para(notes, f'使用 XinYu-V2 当前工作区源码快照，包含未提交修改。共 {summary["source_files"]} 个程序文件，原始物理行数 {summary["original_physical_lines"]:,} 行，其中非空行 {summary["original_nonblank_lines"]:,} 行。统计范围为小程序程序/页面/样式/注册配置、管理后台 src 程序以及 Python 后端 app 和启动脚本；不计第三方依赖、测试、构建产物、锁文件、资源图片和文档。')
para(notes, f'软件名称“{TITLE}”和登记版本号“{VERSION}”为暂用值。目录名 XinYu-V2、代码内部“心语 V2”和包版本 0.1.0/0.0.0 均保持原样；正式提交时需与申请表、说明书统一登记名称和版本。')
para(notes, f'快照时间：{DATA["created_at"]}；Git 基线：{DATA["commit"]}。文件级 SHA-256 和行数见 CSV 清单，完整导出依据见 _build/manifest.json。')
para(notes, '2. 编排与截取规则', bold=True, size=12)
para(notes, f'先按固定文件清单编排全部源码，再从全文 {page_count} 页中提取第1—30页及第{page_count-29}—{page_count}页，共60页。两个片段内部均连续，文件内部没有挑选、拼接或删减业务逻辑。编排顺序为三端入口、学生端页面控制器及其他源文件、后台源文件、后端基础与业务源文件，最后依次为安全支持、树洞和后台工作台服务。')
para(notes, f'每页至少50行非空、非纯注释代码；空白、文件标签、页眉和页脚不计数。共提交 {summary["submitted_counted_code_lines"]} 行计数代码。最后一页到达 admin_workbench_service.py 文件第{summary["final_module_last_line"]}行，完整结束 _write_audit 方法及 AdminWorkbenchService 模块。正文左侧为排版副本文件行号，长字符串仅由 Word 视觉换行。')
para(notes, f'为展开原项目中的超长单行，TypeScript/Vue/页面/样式/JSON 只在导出副本上使用 Prettier {DATA["formatter"].split(";",1)[0].replace("prettier ","")} 排版；Python、Shell 原样保留。已比较 {DATA["transpile_checks"]} 个 TypeScript 文件格式化前后的语法树。排版后行数不用于冒充原始源程序量。原始副本和排版副本同时保存在 _build/ 下，未修改业务源码。')
para(notes, '3. 已执行的运行相关检查', bold=True, size=12)
para(notes, 'Python 3.11.11：后端 pytest 180项通过；小程序 TypeScript 检查和19项测试通过；后台 TypeScript 检查、5项测试和 Vite 构建通过。普通入口在无外部配置时健康接口返回 HTTP 200及 unconfigured / degraded；本地演示入口返回 demo / ok，并已通过实际前端适配器完成自测、安全支持、社区及账户流程核对。OpenAPI注册50个端点。本次前端验证使用现有 Node 26.8.1 / npm 11.19.0；工程规范记录的发布工具链为 Node 22.21.0 / npm 10.9.4。')
para(notes, '上述结论证明本地测试与构建结果。当前代码包含演示数据路径及未配置环境处理；真实微信登录、学校身份核验、CloudBase、DeepSeek及正式部署仍依赖对应环境配置，不能据此写成已完成线上全流程验收。60页为申请用源码摘录，运行项目应使用随附完整工程源代码快照并安装依赖。')
para(notes, '4. 功能模块与代码对应', bold=True, size=12)
modules = [
 ('启动与三端入口', 'miniprogram/app.ts；admin/src/main.ts、App.vue；backend/app/main.py'),
 ('首次使用、身份核验、同意管理', 'miniprogram/pages/onboarding、identity-verification、privacy；backend/app/services/consent_service.py、identity_service.py'),
 ('今日、每日心情与短句', 'miniprogram/pages/today；backend/app/services/today_service.py、mood_service.py、quote_service.py'),
 ('自测、答题、结果与历史', 'miniprogram/pages/assessment-center、assessment-answer、assessment-result、history；backend/app/services/assessment_service.py；backend/app/domain/assessment_rules.py'),
 ('安全确认与支持资源', 'miniprogram/pages/safety-confirmation、support-resources；backend/app/services/safety_service.py、support_resource_service.py'),
 ('树洞浏览、发布、详情及状态', 'miniprogram/pages/treehole、treehole-publish、treehole-detail、treehole-status；backend/app/services/treehole_service.py'),
 ('我的、匿名身份、个人内容与账户操作', 'miniprogram/pages/my、anonymous-identity、my-content、account-stop、account-recover、delete-confirmation；backend/app/services/account_service.py'),
 ('后台任务、审计与演示重置', 'admin/src/views/WorkbenchView.vue、TaskDetailView.vue、AuditDemoView.vue；backend/app/services/admin_workbench_service.py'),
 ('后台身份授权', 'backend/app/services/identity_access_service.py；backend/app/api/admin_identity.py'),
 ('受限 AI 辅助', 'backend/app/services/ai_assist_service.py；backend/app/integrations/deepseek_client.py；backend/app/domain/ai_policy.py'),
]
for name, refs in modules:
    para(notes, name, bold=True)
    para(notes, refs, size=9)
para(notes, '说明书应按已实现代码描述功能，并区分演示状态与实际服务配置状态。上述模块清单包含全量源码中的实现；60页摘录不可能展示每个模块的全部依赖，可使用文件清单追溯。')
para(notes, '5. 提交页码与源文件范围', bold=True, size=12)
for p in manifest_pages:
    refs = '；'.join(f'{r["file"]}:{r["start_line"]}—{r["end_line"]}' for r in p['ranges'])
    para(notes, f'提交第{p["submission_page"]}页 / 全文第{p["full_page"]}页 / {p["code_lines"]}行代码\n{refs}', size=8)
notes.save(OUT / '代码材料核对说明与模块索引.docx')

print(json.dumps({k:v for k,v in summary.items() if k not in ('pages','files')}, ensure_ascii=False, indent=2))
print('WORD:', str(OUT / CODE_NAME))
