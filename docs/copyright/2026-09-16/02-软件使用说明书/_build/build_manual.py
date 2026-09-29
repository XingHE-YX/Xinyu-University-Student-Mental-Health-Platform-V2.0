from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
SHOTS = OUT / '截图'
TITLE = '心语大学生心理健康平台'
VERSION = 'V1.0'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.6), Cm(1.5)
section.left_margin = section.right_margin = Cm(1.8)
section.header_distance = section.footer_distance = Cm(.7)


def set_font(run, size=10.5, bold=False, color='1F2D2B', latin='Arial'):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().rFonts
    if fonts is not None:
        fonts.set(qn('w:eastAsia'), 'Arial Unicode MS')


for name, size in [('Normal',10.5),('Heading 1',17),('Heading 2',13),('Heading 3',11),('Caption',9)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15
    if name.startswith('Heading'):
        style.font.color.rgb = RGBColor.from_string('2E6F68')
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(7)

header = section.header.paragraphs[0]
set_font(header.add_run(f'{TITLE}  {VERSION}'),9,bold=True)
header.add_run('   |   软件使用说明书')
for r in header.runs[1:]: set_font(r,9,color='5F6E6A')
border = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
for k,v in {'val':'single','sz':'4','color':'D7E4DF','space':'5'}.items(): bottom.set(qn('w:'+k),v)
border.append(bottom);header._p.get_or_add_pPr().append(border)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run('软件使用说明书  ·  '),8,color='5F6E6A')
fld = OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

doc.core_properties.title = f'{TITLE} {VERSION} 软件使用说明书'
doc.core_properties.subject = '软件著作权登记配套文档：实际运行截图、技术架构和操作流程'
doc.core_properties.author = ''
doc.core_properties.comments = ''
figures = []


def p(text='', *, bold=False, size=10.5, color='1F2D2B', parent=None):
    para = (parent or doc).add_paragraph()
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.15
    set_font(para.add_run(text),size,bold,color)
    return para


def h(text, level=1, new_page=False):
    para = doc.add_paragraph(text, f'Heading {level}')
    for run in para.runs:set_font(run,17 if level==1 else 13 if level==2 else 11,True,'2E6F68')
    if new_page: para.paragraph_format.page_break_before = True
    return para


def table(headers, rows, widths=None, size=9.5):
    t=doc.add_table(rows=1, cols=len(headers))
    t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(headers):
        cell=t.rows[0].cells[i];cell.text=''
        set_font(cell.paragraphs[0].add_run(x),size,True)
        shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E8F0EC');cell._tc.get_or_add_tcPr().append(shd)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].text='';para=cells[i].paragraphs[0]
            para.paragraph_format.space_after=Pt(3)
            para.paragraph_format.space_before=Pt(3)
            for j,line in enumerate(str(x).split('\n')):
                if j: para.add_run().add_break()
                set_font(para.add_run(line),size)
        no_split=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(no_split)
    if widths:
        t.autofit=False
        for col,w in zip(t.columns,widths):col.width=Cm(w)
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=Cm(w)
    return t


def caption(parent, filename, label, *, interface=True):
    number=len(figures)+1
    title=f'图{number}：{label}' + ('功能界面' if interface else '')
    para=parent.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(4)
    set_font(para.add_run(title),9,color='5F6E6A')
    figures.append({'figure':number,'caption':title,'file':filename})
    return title


def image(parent, filename, label, width, *, diagram=False):
    source=HERE/filename if diagram else SHOTS/filename
    assert source.is_file(), source
    para=parent.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(3)
    para.paragraph_format.keep_with_next=True
    pic=para.add_run().add_picture(str(source),width=Cm(width))
    pic._inline.docPr.set('descr',label)
    caption(parent,filename,label,interface=not diagram)


def pair(title, paragraphs, items, *, source=None):
    h(title,2,new_page=True)
    for x in paragraphs:p(x)
    t=doc.add_table(rows=1,cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c in t.rows[0].cells:
        c.width=Cm(8.7);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        c.paragraphs[0].paragraph_format.space_after=Pt(0)
        c.paragraphs[0].paragraph_format.line_spacing=Pt(1)
        c.paragraphs[0].add_run().font.size=Pt(1)
    for cell,(filename,label) in zip(t.rows[0].cells,items): image(cell,filename,label,7.35)
    no_split=OxmlElement('w:cantSplit');t.rows[0]._tr.get_or_add_trPr().append(no_split)
    if source:p('对应模块：'+source,size=8,color='5F6E6A')


def desktop(title, paragraphs, filename, label, *, source=None, after=None):
    h(title,2,new_page=True)
    for x in paragraphs:p(x)
    image(doc,filename,label,17.1)
    if after:
        for x in after:p(x)
    if source:p('对应模块：'+source,size=8,color='5F6E6A')


# Cover
p('软件著作权登记配套文档',size=11,color='5F6E6A')
p('\n\n')
p(TITLE,bold=True,size=26)
p('软件使用说明书',bold=True,size=23,color='2E6F68')
p(VERSION,size=18,color='5F6E6A')
p('\n')
p('微信小程序学生端 · Python 业务后端 · 桌面 Web 管理后台',size=12)
p('\n\n')
table(['项目','内容'],[
 ('软件名称',TITLE),('登记版本',VERSION),('文档类型','软件使用说明书'),
 ('编制日期','2026年9月16日'),('运行取证','XinYu-V2 当前代码，本地服务与微信开发者工具'),
 ('使用对象','学生用户、心理健康中心工作人员及系统维护人员')
],[3.8,13.5],11)
p('界面中“心语 V2”为项目显示名称；本说明书与源程序材料统一采用上述登记名称和版本。正式申报前应与申请表的最终名称、版本保持一致。',size=10,color='5F6E6A')

h('阅读说明与目录',new_page=True)
p('本说明书依据 XinYu-V2 实际源码和本地运行结果编写。学生端截图来自微信开发者工具，后台截图来自浏览器。操作示例中的姓名、学号、帖子和后台任务均为合成测试数据。截图保留界面标题栏，图号按文档顺序连续编号。')
p('本次取证启动了 Python 后端及 Web 后台，学生端通过本地 HTTP 接口提交和读取业务数据。微信登录和学校身份核验采用本地合成适配器；数据保存在内存中。CloudBase 正式环境和外部 AI 服务未在本次运行中部署或接通。')
p('后台使用项目自带的演示任务展示工作人员操作；这些任务与学生端截图中新发布的帖子属于不同演示记录。本地运行和截图不等同于真实学校环境或云端全流程验收。')
for x in [
 '1  软件概述：开发背景、目标和适用范围',
 '2  功能与角色：学生端和工作人员功能',
 '3  技术实现：架构图、关键技术、数据处理流程',
 '4  运行环境与本地启动',
 '5  学生端操作：首次使用、心情、自测、支持资源、树洞、个人记录与账户',
 '6  管理后台操作：登录、任务处理、审计和状态提示',
 '7  异常处理与使用注意事项',
 '8  后端接口文档',
 '9  功能、界面与源码对应索引',
]:p(x,size=11)
p('图像原文件与说明书一并提供，便于放大检查；文档中的页面状态和按钮名称以截图所示版本为准。',size=10,color='5F6E6A')

h('1  软件概述',new_page=True)
h('1.1  开发背景',2)
p('大学生日常学习与生活中可能出现情绪波动、持续紧张、睡眠节律变化等情况。相关记录常分散在笔记或临时聊天中，缺少统一、低负担的自我观察入口；匿名表达与校园支持资源之间也需要清晰的使用边界。')
p('本软件将每日心情、自测记录、匿名树洞和支持资源整合在微信小程序中，并为工作人员提供独立桌面工作台。学生可以记录与回顾自己的状态，工作人员可以围绕最小必要事实处理任务并回溯操作。')
h('1.2  开发目的与适用范围',2)
p('软件面向在校大学生的日常自我观察、匿名交流和支持资源查询场景，适用于高校心理健康教育与信息服务。它提供记录、固定规则反馈和资源入口，不把自测分数或 AI 文本作为医疗诊断、治疗建议或危机处置结论。')
h('1.3  产品组成',2)
table(['组成','使用入口','主要职责'],[
 ('学生端','微信小程序','今日心情、PHQ-9、GAD-7、睡眠观察、树洞、历史和账户管理'),
 ('业务后端','Python HTTP 服务','认证、同意、身份、固定计分、安全状态、数据操作和审计'),
 ('管理后台','桌面浏览器','任务认领与处理、身份授权任务、事实跟进、只读审计和演示入口'),
],[3.2,4.1,10])
h('1.4  基本使用流程',2)
p('学生通常按“首次说明 → 登录与核验 → 今日 → 心情或自测 → 结果与支持资源 → 树洞或个人记录”使用。工作人员按“登录 → 任务工作台 → 认领 → 查看必要事实 → 提交决定 → 审计回溯”使用。')
p('学生端底部导航固定为“今日、自测、树洞、我的”。答题、发布、确认等页面使用独立操作流程；管理后台不进入小程序底部导航。')

h('2  功能与角色',new_page=True)
table(['角色/模块','主要功能','结果或边界'],[
 ('学生 / 首次使用','阅读基础服务说明、登录、身份核验','真实身份用于核验；树洞展示匿名名'),
 ('学生 / 今日','查看短句、记录当天心情、进入自测和支持资源','心情每日首次保存，当天不提供修改入口'),
 ('学生 / 自测','PHQ-9、GAD-7、睡眠观察','固定规则计算；三个模块不合并总分'),
 ('学生 / 安全支持','PHQ-9第九题触发确认及资源展示','按用户选择进入继续答题、受限结果或支持优先分支'),
 ('学生 / 树洞','浏览匿名帖子、编写提交、查看状态、回应','未公开内容不进入他人的公开列表'),
 ('学生 / 我的','查看匿名身份、本人内容、历史、隐私与账户','逐条记录管理；账户停止后有30天恢复期'),
 ('工作人员 / 任务','查看三个区块，认领、释放和处理任务','内容审核、安全支持、身份授权、跟进四种任务'),
 ('工作人员 / 审计','按条件筛选、分页、查看事件详情','只读展示必要事实；不复制完整答案和身份'),
],[4.0,6.1,7.2])
p('PHQ-9 包含9个计分题和1个不计分的影响题，因此答题页面显示共10题；GAD-7显示7题；睡眠观察显示8题。睡眠观察输出作息节律、睡前阻力、恢复与日间影响三个维度，不输出医学总分。')
p('基础服务同意与社区同意分别管理。撤回社区同意后可以继续浏览公开内容，但不能发布或回应；已有帖子不会因为撤回同意而自动删除。')

h('3  技术实现',new_page=True)
h('3.1  总体架构',2)
image(doc,'architecture.png','软件总体架构图',17.2,diagram=True)
p('系统由三个独立运行单元组成。小程序和 Web 后台经 HTTP API 调用 Python 业务后端；后端承担固定规则、权限与状态校验。外部登录、学校核验、CloudBase 和 DeepSeek 通过明确的适配边界接入。')
h('3.2  关键技术',2)
table(['技术','实现与用途'],[
 ('原生小程序与 TypeScript','使用 Page、Component、WXML、WXSS 和微信原生请求接口；未提交答案保留在当前会话内存'),
 ('Vue 3 与 Vite','实现独立桌面工作台；Pinia保存会话，Vue Router控制页面入口，Zod校验响应边界'),
 ('FastAPI 与 Pydantic','提供 /api/v1 接口，验证请求字段，统一返回 data、error 和 request_id'),
 ('确定性业务规则','服务端校验题卷版本、题目与选项；计算PHQ-9/GAD-7及睡眠维度；安全分支不依赖AI'),
 ('幂等、版本与审计','提交请求使用幂等键；改变已有对象时校验版本；敏感动作记录必要审计事实'),
 ('受限 AI 适配','仅用于自测说明和树洞审核辅助；本次未连接外部模型，固定结果与人工流程独立可用'),
],[4.0,13.3],9.5)

h('3.3  数据处理流程',2,new_page=True)
image(doc,'data-flow.png','业务请求与数据处理流程图',17.2,diagram=True)
p('每项操作先验证会话和输入，再检查账户、同意、身份及对象状态。需要改变数据的动作经过幂等与版本保护后写入记录，并返回当前页面所需的最小字段。失败响应不会被显示为成功。')
p('自测评分只依据服务端发布的题卷和固定计分规则。PHQ-9 第九题非零时先确认当前安全状态；“我不太确定”先展示资源再允许继续；“我现在无法保证安全”进入支持优先页面，不生成完整结果。')
p('后台任务详情采用“必要事实与脱敏内容 / 当前操作”两栏。当前动作受服务端会话、任务状态和对象版本限制；发生冲突时要求重新读取。')

h('4  运行环境与本地启动',new_page=True)
table(['项目','本次验证环境 / 要求'],[
 ('开发与取证主机','macOS；微信开发者工具 Stable 2.01.2510290'),
 ('小程序','原生微信小程序；本次自动化基础库3.16.1，iPhone 12/13 (Pro)模拟器，390px宽'),
 ('后端','Python 3.11.11；FastAPI 0.128.8、Pydantic 2.13.4、httpx 0.28.1、Uvicorn 0.39.0'),
 ('后台','Vue 3.5.42、TypeScript 7.0.2、Vite 8.2.2；桌面浏览器，最小工作宽度1280px'),
 ('前端工具链','项目基线Node 22.21.0/npm 10.9.4；本机验证使用Node 26.8.1/npm 11.19.0'),
 ('本地服务','后端127.0.0.1:9000；后台127.0.0.1:5173；后台/API之间由本地代理转发'),
 ('本地数据','内存仓储和预置合成案例；关闭后端进程后本地新增数据不保留'),
 ('正式运行外部条件','申请人自己的微信AppID、HTTPS服务地址、云端环境、学校核验方式及真实支持资源配置'),
],[4.1,13.2])
h('4.1  启动步骤',2)
for x in [
 '（1）在项目根目录准备后端虚拟环境，并按 backend/requirements.lock 安装依赖；在 admin 和 miniprogram 目录安装各自锁文件中的 npm 依赖。',
 '（2）运行“backend/.venv/bin/python deploy/local/backend_demo.py”，启动仅绑定回环地址的本地演示后端。访问 /api/v1/health，应返回 HTTP 200，环境为 demo，状态为 ok。',
 '（3）另开终端运行“node deploy/local/admin_preview.mjs”，在浏览器打开 http://127.0.0.1:5173。',
 '（4）微信开发者工具打开 miniprogram 目录，使用普通编译。本地预览配置仅在开发者工具且演示模式下接受127.0.0.1地址；正式微信客户端仍要求已配置的HTTPS服务。',
 '（5）本地核验示例填写“演示同学 / DEMO20260916”。后台固定显示“心理健康中心工作人员 / 超级管理员”，本地测试口令见 deploy/local/README.md。',
]:p(x,size=10)
p('本地演示入口是复现说明书界面的辅助配置。正式部署须使用真实服务端环境配置和持久化数据接入，不能以本地合成身份代替学校身份核验。',size=10,color='5F6E6A')

pair('5  学生端操作：首次使用与身份核验',[
 '首次打开后阅读产品边界和基础服务说明，勾选同意并选择“同意并继续”。新用户进入身份核验页，输入姓名、学号并选择“开始核验”；本地示例使用合成测试身份。',
 '核验成功后进入今日页，系统生成匿名展示名。已完成基础同意和身份核验的账户再次登录时可直接回到今日。核验失败时应检查信息或稍后重试。',
],[('01-首次使用.png','首次使用与基础服务同意'),('02-身份核验.png','身份核验')],source='pages/onboarding、identity-verification；services/auth.ts；后端 consent_service、identity_service')

pair('5.2  今日页与快捷入口',[
 '今日页显示日期、问候、短句及出处、当天心情和各类观察入口。没有记录时显示“还没有记录此刻”或“尚未记录”。',
 '页面下部提供树洞和支持资源入口；通过底部导航切换“今日、自测、树洞、我的”。长页面可以上下滚动，微信标题栏和底部主导航仍保留。',
],[('03-今日首页.png','今日首页'),('03B-今日快捷入口.png','今日快捷入口')],source='pages/today；services/today.ts；后端 today_service、quote_service')

pair('5.3  记录当天心情',[
 '点击“记录此刻”，从愉快、平静、疲惫、焦虑、低落、烦躁中选择一项，再点击“保存今天的心情”。截图示例选择“平静”。',
 '成功后显示“已记下此刻”及保存时间。当天只记录一次；再次进入可查看已记录事实。网络失败时保留当前选择并显示重试提示。',
],[('04-选择心情.png','心情选择'),('05-心情保存成功.png','心情保存确认')],source='pages/today；services/mood.ts、today.ts；后端 mood_service')

pair('5.4  自测中心与焦虑自测',[
 '进入“自测”，按需要选择近两周情绪状态、近两周焦虑状态或睡眠与作息观察。本例进入焦虑自测。目录显示用途、预计耗时和最近记录情况。',
 '答题页显示当前题目、进度和固定选项。选择答案后进入下一题；可用“上一题”检查已选答案。离开本次观察前会显示确认提示。',
],[('06-自测中心.png','自测中心'),('07-焦虑自测答题.png','焦虑自测答题')],source='pages/assessment-center、assessment-answer；services/assessment.ts')

pair('5.5  提交自测与查看结果',[
 '完成最后一道题后点击“提交本次观察”。客户端提交题目与选项标识，服务端校验并计算结果；提交处理中不应重复操作。',
 '结果页显示模块、记录时间、本次分数、固定说明和下一步入口。示例的GAD-7七题均选择第一项，服务端返回0分。用户可继续查看支持资源、历史记录或再次观察。',
],[('08-焦虑自测提交.png','自测完成提交'),('09-焦虑自测结果.png','焦虑自测结果')],source='assessment_service.py、assessment_rules.py；pages/assessment-result')

pair('5.6  睡眠观察',[
 '从自测中心进入睡眠与作息观察，按最近7天的实际情况完成8道题。题目及选项来自服务端题卷。',
 '结果显示作息节律、睡前阻力、恢复与日间影响三个维度。本模块不计算医学总分。向下滚动可查看支持资源、历史记录和再次观察入口。',
],[('11-睡眠观察答题.png','睡眠观察答题'),('12-睡眠观察结果.png','睡眠观察结果')],source='assessment_rules.py 的睡眠规则；services/assessment.ts 的维度映射')

pair('5.7  情绪自测与安全确认触发',[
 '进入情绪自测，依次完成题目。页面共10项，其中9项用于计分，最后一项为影响题。截图显示第1项和第9项。',
 '当第9项选择非零选项时，立即暂停普通答题并进入安全确认；系统不等待AI判断。示例仅用于展示软件分支，不代表真实用户状态。',
],[('13-情绪自测答题.png','情绪自测答题'),('14-情绪自测第九题.png','第九题安全分支触发')],source='pages/assessment-answer；后端 safety_rules、safety_service')

pair('5.8  安全确认与支持资源',[
 '安全确认提供“我现在可以保证安全”“我不太确定”“我现在无法保证安全”三项选择。选择可以保证安全时返回继续答题；不太确定时先展示支持资源；无法保证安全时优先展示支持资源，不进入完整结果。',
 '本例选择“我不太确定”。资源页面显示安全提示及当前环境配置的资源；阅读后可点击“返回继续观察”。本地资源明确标注为演示内容，不提供真实服务承诺。',
],[('15-安全确认.png','安全确认'),('16-安全支持资源.png','安全支持资源')],source='pages/safety-confirmation、support-resources；safety_service、support_resource_service')

pair('5.9  受限结果与普通支持资源',[
 '从“不太确定”分支继续并完成剩余题目后，结果仅保存和展示必要支持事实，不展示分数、AI内容或历史入口。截图中的状态是完成记录的事实，不表示用户已经安全。',
 '普通场景也可以从今日、结果或我的页面进入支持资源，查看资源描述与更新时间。有真实电话配置时按页面提供的拨号入口操作；缺失资源时页面提示暂未配置。',
],[('17-受限安全结果.png','受限安全结果'),('10-支持资源.png','普通支持资源')],source='pages/assessment-result、support-resources；AssessmentResultProjection')

pair('5.10  浏览树洞与社区同意',[
 '进入“树洞”可浏览允许公开的匿名帖子。列表显示匿名名、时间和可展示摘要；点击“读一读”进入详情。',
 '首次发布或社区同意已撤回时，点击“写下此刻”进入数据与隐私页面。阅读社区说明并选择“同意参与社区”，再返回树洞继续发布。',
],[('18-树洞首页.png','匿名树洞列表'),('19-社区同意与隐私.png','社区内容同意')],source='pages/treehole、privacy；treehole_service、consent_service')

pair('5.11  编辑、检查与确认提交树洞',[
 '在“写一段树洞”页面输入正文，留意1000字长度提示。点击“继续检查”后阅读检查提示；格式检查通过不等于内容已经公开。',
 '确认内容后选择“确认提交”。请求成功后进入帖子状态页，后续可在“我的树洞内容”查看本人帖子状态。示例正文明确标注为本地演示内容。',
],[('20-树洞内容编辑.png','树洞内容编辑'),('21-树洞提交确认.png','树洞提交确认')],source='pages/treehole-publish；services/treehole.ts；treehole_service')

pair('5.12  帖子状态、详情与回应',[
 '提交后应以服务端状态为准。本例显示“检查中”，表示检查完成前不会出现在公开列表，不能把提交成功理解为已公开。',
 '公开帖子详情显示允许展示的正文、匿名身份和回应区域。社区同意有效时可以输入并提交回应；提交后刷新可公开的回应列表，尚在检查中的回应不作为公开内容显示。',
],[('22-树洞提交状态.png','帖子提交状态'),('23-树洞详情与回应.png','帖子详情与回应')],source='pages/treehole-status、treehole-detail；TreeholePostProjection')

pair('5.13  我的页面与匿名身份',[
 '点击底部“我的”，可进入匿名身份、本人树洞内容、历史记录、支持资源、数据与隐私以及账户操作。首页不公开真实姓名和学号。',
 '“在树洞中的身份”展示服务端生成的匿名名及用途说明。匿名身份用于树洞展示，不由用户填写真实身份作为公开昵称。',
],[('24-我的首页.png','个人中心'),('24B-匿名身份.png','匿名身份')],source='pages/my、anonymous-identity；/app/bootstrap、/me/anonymous-identity')

pair('5.14  本人内容与历史记录',[
 '“我的树洞内容”列出本人发布的内容及当前状态。检查中或不允许展示正文的内容只显示必要状态，不显示后台内部理由。',
 '“我的记录”按时间倒序显示心情、自测和睡眠记录。截图中的phq9、gad7和sleep_observation分别表示情绪自测、焦虑自测和睡眠观察；时间字符串带Z时表示UTC时间。',
],[('24C-我的树洞内容.png','本人树洞内容'),('25-历史记录.png','本人历史记录')],source='pages/my-content、history；services/me.ts、treehole.ts')

pair('5.15  单条删除与社区同意撤回',[
 '在历史列表选择一条记录进入“删除记录”确认页。确认页说明删除范围与不可恢复性；选择“确认删除”提交该条记录，选择“暂不删除”返回。本次截图保留确认流程中的示例记录。',
 '进入数据与隐私，可撤回社区内容同意。撤回后按钮变为“同意参与社区”，浏览公开树洞仍可用，发帖和回应需重新同意；已有内容不会自动删除。',
],[('26-单条记录删除确认.png','单条记录删除确认'),('27-撤回社区同意.png','社区同意撤回')],source='pages/delete-confirmation、privacy；account/consent 相关服务')

pair('5.16  停止使用与恢复账户',[
 '在我的或数据与隐私中选择“停止使用账户”，阅读影响范围，输入确认词“停止使用”，再选择“确认停止使用”。',
 '成功后账户进入30天恢复期，并显示恢复入口。恢复期内选择“恢复账户”可重新启用账户并回到今日；已删除记录不会随恢复自动还原。本地已实际验证停止与恢复操作。',
],[('28-停止使用账户.png','停止使用账户确认'),('29-账户恢复入口.png','恢复账户入口')],source='pages/account-stop、account-recover；account_service')

h('5.17  恢复后检查',2,new_page=True)
p('恢复后进入“我的”，检查账户入口和身份摘要。恢复期说明及“恢复账户”入口消失，正常显示“停止使用账户”，表示账户已回到可用状态。此前撤回的社区同意仍保持撤回，不会因账户恢复自动重新同意。')
image(doc,'29B-账户恢复后.png','账户恢复后的个人中心',7.35)

desktop('6  管理后台操作：工作人员登录',[
 '使用桌面浏览器访问后台地址。页面显示固定账号“心理健康中心工作人员”和能力“超级管理员”，输入当前环境配置的口令并点击“进入后台”。',
 '登录成功后进入任务工作台。界面显示演示或授权模式及会话有效期；低于1280px宽度时提示使用电脑浏览器。',
],'30-后台登录.png','工作人员登录',source='AdminLoginView.vue；adminAuth.ts；/admin/auth/login')

desktop('6.2  三任务区块与筛选',[
 '工作台分为“需要我处理、等待他人、最近处理”三个区块。可按任务类型、状态和更新时间筛选；各区块分别展示任务摘要与进入详情的按钮。',
 '从“需要我处理”选择一项进入详情。“等待他人”用于查看当前不可直接处理的任务；“最近处理”用于回看已发生的操作事实。',
],'31-后台工作台.png','统一任务工作台',source='WorkbenchView.vue、TaskSection.vue、TaskCard.vue；admin_workbench_service')

desktop('6.3  查看必要事实并认领任务',[
 '任务详情左侧展示来源、状态和脱敏内容，右侧展示当前操作。未认领时只提供“开始处理”；先检查必要事实，再进行认领。',
 '本例为项目内置的内容审核演示任务。认领成功后状态变为“已认领”，页面显示当前任务允许执行的决定选项。',
],'32-内容审核待认领.png','内容审核任务详情',source='TaskDetailView.vue、TaskFactsPanel.vue；/admin/tasks/{task_id}/claim')

desktop('6.4  填写审核决定',[
 '认领后可以选择公开、保护展示、暂不公开或转安全复核。本例选择“保护展示”，并在内部事实理由中填写简短的处理依据。',
 '内部理由用于任务记录和审计，不展示给学生。填写时应记录实际处理依据，不写诊断、身份推断或与当前任务无关的信息。',
],'33-审核决定.png','内容审核决定',source='TaskActionPanel.vue；admin_workbench_service.decision')

desktop('6.5  提交前确认',[
 '点击“提交决定”后，页面显示“提交前确认”，再次列出将执行的动作和审计范围。确认无误后选择“确认提交”，需要调整时选择“返回修改”。',
 '系统在提交时重新检查权限、对象版本和任务状态。任务已变化时应重新读取，不用旧页面覆盖其他操作。',
],'34-审核提交确认.png','审核提交确认',source='TaskDetailView.vue；/admin/tasks/{task_id}/decision')

desktop('6.6  查看处理完成状态',[
 '成功后任务状态显示“已完成”，左侧新增已发生的事实记录，右侧提示当前任务仅供查看。返回工作台后可在最近处理区域回看。',
 '本地演示记录验证了任务状态与审计写入。对真实帖子公开状态的同步，需要正式环境将工作台服务与对应内容仓储完成接入。',
],'35-审核处理完成.png','任务处理完成',source='admin_workbench_service；TaskStateBar.vue')

desktop('6.7  安全支持任务',[
 '安全支持任务只展示最小必要事实和当前状态。认领后可选择确认支持已提供、记录一次事实跟进或结束本次支持记录，并按需要记录下一次跟进时间。',
 '只记录已发生的支持事实。演示环境中的操作不表示已经联系真实学生，也不表示用户已处于安全状态。',
],'37-后台安全支持.png','后台安全支持任务',source='TaskActionPanel.vue 的 safety_support 变体')

desktop('6.8  身份授权任务',[
 '身份授权详情显示申请范围、有效期及当前状态。认领后可以批准、拒绝或撤回授权；服务端仍需检查授权范围和对象版本。',
 '截图是合成任务的操作页面，不包含真实姓名、学号或身份读取结果。正式使用时按申请范围读取必要字段，并保留相应审计。',
],'38-后台身份授权.png','后台身份授权任务',source='TaskActionPanel.vue；identity_access_service、admin_identity.py')

desktop('6.9  记录跟进事实',[
 '跟进任务提供“保存跟进记录”与“结束本次支持记录”。选择记录类型，填写事实说明，必要时设置下一次跟进时间，再经过提交确认。',
 '本例输入“本地演示：已核对支持资源入口，未进行真实联系”。演示模式不提供把测试操作标记成真实联系已完成的选项。',
],'40-后台跟进记录.png','后台跟进记录',source='TaskActionPanel.vue 的 followup 变体')

desktop('6.10  审计筛选、分页与详情',[
 '进入“审计与演示”，可按时间、事件类型、结果、模式和对象编号筛选。列表支持分页，选中事件后查看操作者、动作、时间、字段范围及请求编号。',
 '审计页面只读，不提供修改或删除日志的操作。内容只保存必要事实，不复制完整答题、完整树洞正文或不必要的身份信息。',
],'36-审计日志.png','只读审计日志',source='AuditDemoView.vue、AuditTable.vue；/admin/audit-events')

desktop('6.11  会话失效处理',[
 '会话过期时页面提示“登录状态已过期，请重新登录”，停止当前受保护操作并提供“返回登录”。重新登录后再读取任务，避免使用过期详情。',
 '截图来自本地实际发生的后台会话到期。并发冲突、无权限和提交失败也会在详情区域显示对应提示；应按提示重新读取或重试。',
],'39-会话失效.png','后台会话失效',source='SessionExpiredState.vue、ConcurrencyState.vue、PermissionState.vue')

desktop('6.12  演示重置入口',[
 '在演示环境的审计页面点击“重置演示数据”，系统先显示环境、数据范围和二次确认。选择取消可返回审计页面。真实环境应由服务端拒绝演示重置。',
 '本地当前实现恢复内置后台任务案例并写入审计；不能据此认定已经清理正式CloudBase中的全部数据域。本次保留确认页面作为操作说明，未执行最终重置。',
],'41-演示重置确认.png','演示数据重置确认',source='DemoResetDialog.vue；admin_workbench_service.reset_demo')

h('7  异常处理与使用注意事项',new_page=True)
table(['页面提示 / 情况','处理方式','系统行为'],[
 ('网络暂时不可用','检查连接，使用页面重试入口','未成功写入时不显示保存成功'),
 ('会话已过期','返回首次使用或后台登录，重新建立会话','重新读取本人信息或任务详情'),
 ('身份核验未通过','核对输入；服务不可用时稍后重试','不把失败或不可用状态显示为已核验'),
 ('未同意参与社区','到数据与隐私阅读并完成社区同意','未同意时禁止发布与回应'),
 ('内容刚刚发生变化','重新加载并依据最新对象状态操作','避免旧版本覆盖新的处理结果'),
 ('内容检查中','到我的内容查看当前状态','检查中内容不进入公开列表'),
 ('支持资源未配置','根据页面提示联系可信任的人或通过真实公开渠道查询','不生成虚构联系方式'),
 ('账户处于恢复期','在允许期限内进入恢复账户','恢复前不接受新的记录、帖子或回应'),
 ('AI辅助不可用','使用固定规则结果与现有支持资源','不影响固定计分，不由AI决定安全状态'),
],[4.8,6.0,6.5])
p('账户记录、自测和后台任务应以服务端返回状态为准。软件的自测用于自我观察，结果不是诊断。截图中的演示资源、演示身份和演示任务不能直接用于真实学校业务。')
p('本地截图中的部分时间字段直接显示服务端UTC值；带“Z”的时间为UTC，微信状态栏显示本机时间。正式部署时应统一展示时区及具体支持资源信息。')

h('8  后端接口文档',new_page=True)
h('8.1  接口约定',2)
p('后端不单独提供业务操作界面，学生端和后台经 /api/v1 调用。以下接口清单从本次启动服务的 OpenAPI 导出，实际注册50个方法与路径组合；完整机器可读结构随材料保存在 _build/openapi.json。')
table(['约定','说明'],[
 ('传输','正式环境HTTPS；本地演示使用回环HTTP'),
 ('数据格式','请求与响应为JSON；有请求体时使用 Content-Type: application/json'),
 ('认证','除健康检查与认证入口外，业务接口按需要校验 Authorization: Bearer <access_token>'),
 ('请求追踪','X-Request-ID；响应 request_id 用于定位该次请求'),
 ('幂等','需要的提交类接口提供 Idempotency-Key；部分创建接口在请求体中提供 client_start_key 或 client_idempotency_key'),
 ('并发控制','修改已有对象时提交服务端返回的 object_version；冲突后重新读取，不能固定写成1'),
 ('统一响应','成功：data为结果、error为null；失败：error包含稳定错误码、消息与可重试标志'),
],[4.0,13.3])
h('8.2  主要业务数据',2)
table(['数据对象','主要字段','用途'],[
 ('题卷与会话','module_code、questionnaire_version、session_id、questions、object_version','冻结题卷，按真实question_key与option_key提交答案'),
 ('自测完成','answers[]、object_version → completion_state、result_id','完成后用result_id读取完整结果投影'),
 ('安全确认','state、answers[]、object_version → next_step、version','保存确认状态，并用新版本确认已展示资源'),
 ('资源确认','resource_context、resource_version、object_version','记录实际展示资源版本；不等于用户已安全'),
 ('树洞内容','body、client_idempotency_key、visibility_state','发布并返回允许展示的投影'),
 ('后台任务','task_id、object_version、action、事实字段','认领、释放或记录处理决定'),
],[3.2,8.0,6.1],9)

spec=json.loads((HERE/'openapi.json').read_text())
schemas=spec.get('components',{}).get('schemas',{})
def resolve(schema):
    if '$ref' in schema:return schemas[schema['$ref'].rsplit('/',1)[-1]]
    return schema
purpose = {
 'health':'健康检查','wechat_session':'微信凭证换学生会话','refresh':'刷新学生会话','logout':'退出学生会话','me':'读取学生主体与账户状态',
 'login':'后台固定账号登录','admin_login':'后台固定账号登录','admin_refresh':'刷新后台会话','admin_logout':'退出后台会话','admin_me':'读取后台账号摘要',
 'bootstrap':'读取账户、同意与匿名身份摘要','base_consent':'记录基础服务同意','community_consent':'记录或撤回社区同意',
 'verify_identity':'提交身份核验','verification_status':'读取身份核验状态','anonymous_identity':'读取匿名身份','today':'读取今日内容',
 'put_mood':'保存今日心情','moods':'读取本人心情历史','delete_mood':'删除本人单条心情','support_resources':'读取支持资源',
 'stop_account':'停止使用账户','recover_account':'恢复账户','account_status':'读取账户及恢复状态',
 'assessment_modules':'读取自测目录','start_assessment_session':'开始自测会话','complete_assessment_session':'完成自测并返回结果引用',
 'abandon_assessment_session':'放弃未完成会话','assessment_result':'读取单条结果','assessment_results':'读取本人结果列表',
 'delete_assessment_result':'删除本人结果','safety_confirmation':'确认安全状态','support_resource_ack':'确认已展示资源版本',
 'workbench':'读取后台任务区块','task_detail':'读取后台任务详情','claim_task':'认领任务','release_task':'释放任务','decide_task':'提交任务处理决定',
 'audit_events':'查询审计日志','reset_demo':'重置后台演示案例','create_identity_access_request':'创建身份授权申请','get_identity_access_request':'读取身份授权申请','read_identity':'按授权范围读取身份字段',
 'list_posts':'读取公开树洞列表','create_post':'提交树洞帖子','get_post':'读取帖子允许展示的详情','delete_post':'删除本人帖子','withdraw_post':'撤回本人帖子','create_response':'提交帖子回应','delete_response':'删除本人回应','list_my_posts':'读取本人帖子列表',
}
endpoint_records=[]
for route,methods in spec['paths'].items():
    for method,operation in methods.items():
        if method not in {'get','post','put','patch','delete'}:continue
        opid=operation.get('operationId','')
        name=opid.split('_api_v1')[0]
        label=purpose.get(name,operation.get('summary',name).replace('_',' '))
        body=operation.get('requestBody',{}).get('content',{}).get('application/json',{}).get('schema',{})
        schema=resolve(body)
        required=set(schema.get('required',[]))
        body_fields=[k+('（必填）' if k in required else '') for k in schema.get('properties',{})]
        params=[q['name']+'['+q['in']+']' for q in operation.get('parameters',[]) if q['name'].lower()!='authorization']
        response=operation.get('responses',{}).get('200',{}).get('content',{}).get('application/json',{}).get('schema',{})
        response_name=response.get('$ref','').rsplit('/',1)[-1] or 'JSON响应'
        endpoint_records.append({'method':method.upper(),'path':route,'purpose':label,'parameters':params,'body_fields':body_fields,'response':response_name})

h('8.3  全部接口清单',2,new_page=True)
p('路径中的花括号为路径参数；“请求字段”列为请求体字段，标记必填的字段必须提供。具体枚举、长度和嵌套结构以随附OpenAPI为准。',size=9.5)
for i,e in enumerate(endpoint_records,1):
    start_index=len(doc.paragraphs)
    para=p(f'{i:02d}  {e["method"]}  {e["path"]}',bold=True,size=9.2)
    para.paragraph_format.keep_with_next=True
    p('用途：'+e['purpose']+'。'+(' 参数：'+'、'.join(e['parameters'])+'。' if e['parameters'] else ''),size=9)
    if e['body_fields']:p('请求字段：'+'、'.join(e['body_fields'])+'。',size=8.5)
    p('成功响应：'+e['response']+'，外层包含request_id、data、error。',size=8.3,color='5F6E6A')
    block=doc.paragraphs[start_index:]
    for para in block:
        para.paragraph_format.space_after=Pt(2)
        para.paragraph_format.line_spacing=1.05
        para.paragraph_format.keep_with_next=True
    block[-1].paragraph_format.keep_with_next=False
    block[-1].paragraph_format.space_after=Pt(7)

h('8.4  错误响应与处理',2,new_page=True)
table(['HTTP状态','典型含义','客户端处理'],[
 ('400','缺少幂等标识或请求格式不符合约定','按当前接口要求补全请求'),
 ('401','未登录或会话失效','重新登录，再读取需要保护的内容'),
 ('403','权限、同意或身份前置条件不满足','完成前置流程或停止当前动作'),
 ('404','对象不存在或不可见','返回列表并重新读取'),
 ('409','对象版本或幂等请求冲突','读取最新对象，不覆盖新状态'),
 ('422','字段、题目、选项或确认词无效','根据接口约束修正输入'),
 ('503','外部依赖未配置或暂不可用','显示不可用状态并允许适当重试'),
],[2.7,7.2,7.4])
p('自测客户端只提交题目和选项标识，不提交自算分数、安全结论或公开状态。后端验证通过后返回保存结果，客户端再读取完整结果内容。')
p('安全确认返回的新version需要用于后续资源确认和完成提交；账户停止与恢复使用读取到的最新object_version。这样可以避免固定版本值引起错误的并发冲突。')

h('9  功能、界面与源码对应索引',new_page=True)
table(['功能','界面目录 / 文件','核心服务'],[
 ('首次使用与身份','miniprogram/pages/onboarding、identity-verification','auth.ts；consent_service.py；identity_service.py'),
 ('今日与心情','miniprogram/pages/today','today.ts；mood_service.py；today_service.py'),
 ('自测目录与答题','pages/assessment-center、assessment-answer','assessment.ts；assessment_service.py；assessment_rules.py'),
 ('结果与安全支持','pages/assessment-result、safety-confirmation、support-resources','safety_service.py；support_resource_service.py'),
 ('匿名树洞','pages/treehole、treehole-publish、treehole-detail、treehole-status','treehole.ts；treehole_service.py'),
 ('个人信息与记录','pages/my、anonymous-identity、my-content、history、delete-confirmation','me.ts；auth.ts；本人数据投影接口'),
 ('隐私与账户','pages/privacy、account-stop、account-recover','consent_service.py；account_service.py'),
 ('后台登录与导航','admin/src/views/AdminLoginView.vue；router/index.ts','adminAuth.ts；auth_service.py'),
 ('后台任务','WorkbenchView.vue；TaskDetailView.vue；TaskActionPanel.vue','admin_workbench_service.py'),
 ('审计与状态','AuditDemoView.vue；AuditTable.vue；各状态组件','audit/writer.py；admin_workbench.py'),
 ('外部适配','backend/app/integrations；repositories/cloudbase_gateway.py','微信、学校核验、DeepSeek与CloudBase适配边界'),
],[3.0,7.2,7.1],8.6)
p('本说明书与同批次更新的源程序材料对应。源程序材料按完整程序的固定文件顺序提取连续前30页、后30页，故部分完整实现位于未摘录的中间段；可结合源文件清单与完整工程快照追溯。',size=10)
p('截图说明：所有截图为2026年9月16日本地运行界面的直接捕获。后台演示任务、支持资源示例及测试身份均为合成数据。后端接口定义由本地启动后的OpenAPI导出。',size=10,color='5F6E6A')

file=OUT/f'{TITLE}_{VERSION}_软件使用说明书.docx'
doc.save(file)
(HERE/'figure-index.json').write_text(json.dumps(figures,ensure_ascii=False,indent=2))
(HERE/'endpoint-index.json').write_text(json.dumps(endpoint_records,ensure_ascii=False,indent=2))
print(json.dumps({'file':str(file),'figures':len(figures),'screenshots':sum(not x['file'].endswith(('architecture.png','data-flow.png')) for x in figures),'endpoints':len(endpoint_records)},ensure_ascii=False,indent=2))
