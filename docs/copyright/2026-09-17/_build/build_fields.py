from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

out=Path(__file__).resolve().parents[1]
fields=[
('软件开发硬件环境','Apple M5，16GB内存，512GB SSD',50),
('软件运行硬件环境','服务端：Apple M5、16GB内存；客户端：智能手机或电脑',50),
('开发该软件的操作系统','macOS 26',50),
('软件开发环境/工具','微信开发者工具、Vite 8.2.2、Vue 3.5.42、Node.js 26.8.1',50),
('软件运行平台/操作系统','服务端macOS 26；学生端微信小程序；管理端桌面浏览器',50),
('软件运行支撑环境/支持软件','Python 3.11.11、FastAPI 0.128.8、Uvicorn 0.39.0',None),
('编程语言','Python 3.11.11、TypeScript 7.0.2；WXML、WXSS、HTML、CSS',120),
('源程序量（行）','16753行',None)]
for _,value,limit in fields:
 if limit:assert len(value)<=limit

doc=Document();s=doc.sections[0]
s.page_width,s.page_height=Cm(21),Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.7)
s.left_margin=s.right_margin=Cm(1.6)
s.header_distance=s.footer_distance=Cm(.7)

def font(r,size=10.5,bold=False,color='1F2D2B'):
 r.font.name='Arial';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')

def p(text,size=10.5,bold=False):
 t=doc.add_paragraph();t.paragraph_format.space_after=Pt(7);t.paragraph_format.line_spacing=1.15
 font(t.add_run(text),size,bold);return t

p('软著申请表技术信息修正版',18,True)
p('对应：心语大学生心理健康平台 V1.0｜核对日期：2026-09-17',10)
p('以下8项可直接替换截图中的对应内容。填写口径与2026年9月16日交付的源码材料及本地验证说明书保持一致。',10)
t=doc.add_table(rows=1,cols=3);t.style='Table Grid';t.autofit=False
widths=[4.6,11.1,2.0]
for col,w in zip(t.columns,widths):col.width=Cm(w)
for c,txt in zip(t.rows[0].cells,['表格字段','修正后的填写内容','字符数']):
 font(c.paragraphs[0].add_run(txt),10.5,True)
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8F0EC');c._tc.get_or_add_tcPr().append(sh)
for name,value,limit in fields:
 row=t.add_row()
 for c,w,txt in zip(row.cells,widths,[name,value,str(len(value))+(f'/{limit}' if limit else '')]):
  c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  pp=c.paragraphs[0];pp.paragraph_format.space_before=pp.paragraph_format.space_after=Pt(5)
  pp.paragraph_format.line_spacing=1.1;font(pp.add_run(txt),10)
 no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
p('核对说明',12,True)
notes=[
'1. Node.ja应为Node.js；Vue.5.42应为Vue 3.5.42。开发工具栏采用说明书本次验证的Node.js 26.8.1；项目规范基线22.21.0本身不是错误，应区分规范基线与实际验证版本。',
'2. 390px、1280px是界面尺寸；iPhone 12/13 (Pro)是此次模拟器型号；本地回环和HTTPS是网络条件。它们不宜混入运行硬件栏，本版已归并为硬件、平台及支撑软件三类信息。',
'3. 当前主机为macOS 26.7，截图写26.5.1；开发操作系统应按实际开发时记录。这里采用共同大版本macOS 26，未把当前补丁号当成历史开发版本。',
'4. 已交付源码快照为16,753物理行，不是2.1万行；排版后行数及第三方依赖不能代替该统计值。2026-09-17按同一文件清单读取当前代码为16,778行。若改用最新代码申报，应同时更新源程序材料和对应行数。',
'5. 编程语言保留Python和TypeScript的实际版本，并补充源码中的WXML、WXSS、HTML、CSS；不为无统一独立版本的标记或样式语言编造版本号。',
]
for n in notes:p(n,9)
p('字符数包含标点、空格及英文字母；标为50字符、120字符的项目均已校验。原始申请表文件仍是空白模板，本修正版仅整理截图涉及的技术信息。',9)
footer=s.footer.paragraphs[0];font(footer.add_run('技术信息核对稿 · 与同批次软著材料统一使用'),8,color='5F6E6A')
doc.core_properties.title='软著申请表技术信息修正版';doc.core_properties.author=''
filename=out/'软著申请表技术信息修正版.docx';doc.save(filename)
(out/'技术字段修正.json').write_text(json.dumps([{'field':a,'value':b,'characters':len(b),'limit':c} for a,b,c in fields],ensure_ascii=False,indent=2))
print(filename)
print([(a,len(b),c) for a,b,c in fields])
